from pathlib import Path
from typing import Dict, List
from docx import Document

class WordReportBuilder:
    def __init__(self, output_path: Path):
        self.output_path = output_path
        self.document = self._load_or_create()

    def _load_or_create(self) -> Document:
        if self.output_path.exists():
            return Document(self.output_path)
        doc = Document()
        doc.add_heading("Claim Processing Report", level=1)
        return doc

    # ===============================
    # SECTION 1: CLAIM METADATA
    # ===============================
    def add_claim_metadata(self, metadata: Dict):
        self.document.add_heading("CASE INFORMATION", level=2)
        table = self.document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key, value in metadata.items():
            row = table.add_row().cells
            row[0].text = str(key)
            clean_value = str(value).strip() if value is not None else ""
            row[1].text = clean_value if clean_value else "---"

    # ===============================
    # SECTION 2: LINE ITEMS
    # ===============================
    def add_line_items(self, rows: List[Dict], columns: List[Dict]):
        self.document.add_heading("Line Items", level=2)
        table = self.document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        # Header
        for i, col in enumerate(columns):
            table.rows[0].cells[i].text = col["label"]
        # Data
        for row_data in rows:
            cells = table.add_row().cells
            for i, col in enumerate(columns):
                cells[i].text = str(row_data.get(col["key"], ""))

    def save(self):
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        self.document.save(self.output_path)
